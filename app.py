import os
import io
import time
import pandas as pd
import streamlit as st
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# -----------------------------------------------------------------------------
# 1. CONFIGURAÇÃO DA PÁGINA DO STREAMLIT
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Piloto IA CIN - Gestão de RFAs",
    page_icon="🏗️",
    layout="wide",
)

st.title("🏗️ Piloto IA CIN — Automação de RFAs e Investimentos")
st.caption("Coordenadoria de Infraestrutura Educacional (SESI/SENAI SC) — Sistema FIESC")

# Barra Lateral (Sidebar)
with st.sidebar:
    st.header("⚙️ Configurações & Status")
    st.info("Padrão SE Suíte Ativo\nMapeamento Contábil CAPEX/OPEX Habilitado")
    st.divider()
    st.caption("Desenvolvido para automação da emissão de minutas de RFA.")

# -----------------------------------------------------------------------------
# 2. DICIONÁRIO DE SINÓNIMOS E PROCESSAMENTO DE DADOS
# -----------------------------------------------------------------------------
DICIONARIO_FATORES = {
    'AMBIENTE': ['Ambiente', 'Espaço', 'Sala', 'Local', 'Laboratório', 'Andar', 'Bloco'],
    'ITEM': ['Item', 'Descrição', 'Equipamento', 'Descrição do Item', 'Especificação', 'Conta', 'CONTA'],
    'QUANTIDADE': ['Qtd', 'Quantidade', 'Qtd.', 'Quant', 'Unidades'],
    'VALOR_UNIT': ['Valor Unitário Previsto', 'Valor Unitário', 'Valor Unit', 'Preço Unit'],
    'VALOR_TOTAL': ['Valor total Previsto', 'Valor Total', 'Total Previsto', 'Valor Total Final', 'VALOR_TOTAL'],
    'TIPO_DESPESA': ['Tipo de Despesa', 'Classificação', 'Natureza', 'CAPEX/OPEX', 'TIPO_DESPESA'],
    'GETIC_CHAMADO': ['Chamado GETIC', 'Ticket GETIC', 'GETIC', 'Chamado TI'],
    'GENGE_CHAMADO': ['Chamado GENGE', 'Ticket GENGE', 'GENGE', 'Chamado Eng', 'CHAMADO']
}

def processar_planilha(file):
    df = pd.read_excel(file)
    df.columns = df.columns.astype(str).str.strip()
    
    mapeamento = {}
    for chave, sinonimos in DICIONARIO_FATORES.items():
        for col in df.columns:
            for sinonimo in sinonimos:
                if sinonimo.upper() in col.upper() and col not in mapeamento.values():
                    mapeamento[chave] = col
                    break
            if chave in mapeamento:
                break
                
    df_padronizado = df.rename(columns={v: k for k, v in mapeamento.items()})
    
    # Garantia do cálculo do valor total por item
    if 'VALOR_TOTAL' not in df_padronizado.columns or df_padronizado['VALOR_TOTAL'].sum() == 0:
        if 'QUANTIDADE' in df_padronizado.columns and 'VALOR_UNIT' in df_padronizado.columns:
            df_padronizado['VALOR_TOTAL'] = df_padronizado['QUANTIDADE'] * df_padronizado['VALOR_UNIT']
        else:
            df_padronizado['VALOR_TOTAL'] = 0.0
            
    if 'TIPO_DESPESA' not in df_padronizado.columns:
        df_padronizado['TIPO_DESPESA'] = 'CAPEX'
        
    return df_padronizado

# -----------------------------------------------------------------------------
# 3. FUNÇÕES AUXILIARES DE FORMATAÇÃO PARA WORD (DOCX)
# -----------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def gerar_docx_rfa_bytes(gerente_resp, unidade_regional, descricao_sucinta, periodo_desembolso, 
                         justificativa_gestor, observacoes_usuario, resumo_contas, 
                         total_capex, total_opex, total_geral, chamados_getic_str, chamados_genge_str):
    doc = Document()

    # Configuração de Margens (2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Estilo de Fonte
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Cabeçalho
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("SISTEMA FIESC / SENAI SC — COORDENADORIA DE INFRAESTRUTURA")
    run_title.bold = True
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(0x00, 0x56, 0xA6)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("REQUISIÇÃO DE AUTORIZAÇÃO DE DESPESA / INVESTIMENTO (RFA)")
    run_sub.bold = True
    run_sub.font.size = Pt(14)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. Identificação
    h1 = doc.add_paragraph()
    r1 = h1.add_run("1. IDENTIFICAÇÃO E SOLICITAÇÃO")
    r1.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(0x00, 0x56, 0xA6)

    dados_solicitacao = [
        ("Gerente Responsável:", gerente_resp),
        ("Descrição Sucinta:", descricao_sucinta),
        ("Valor Total Previsto:", f"R$ {total_geral:,.2f}"),
        ("Período de Desembolso:", periodo_desembolso),
        ("Regional / Unidade:", unidade_regional),
        ("Linha de Investimento:", "SENAI - Tecnologia, equipamentos para recheios Escolas"),
        ("Tipo de RFA:", "Investimento"),
        ("Despesas Orçadas?:", "Não"),
        ("Vínculos de Governança:", f"GETIC: {chamados_getic_str} | GENGE: {chamados_genge_str}")
    ]

    table_info = doc.add_table(rows=len(dados_solicitacao), cols=2)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, (rotulo, valor) in enumerate(dados_solicitacao):
        row = table_info.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = Inches(2.2), Inches(4.5)
        cell_lbl.paragraphs[0].add_run(rotulo).bold = True
        cell_val.paragraphs[0].add_run(str(valor))
        set_cell_background(cell_lbl, 'F2F4F7')
        set_cell_margins(cell_lbl, top=60, bottom=60)
        set_cell_margins(cell_val, top=60, bottom=60)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 2. Justificativa
    h2 = doc.add_paragraph()
    r2 = h2.add_run("2. JUSTIFICATIVA TÉCNICA E ECONÔMICA")
    r2.bold = True
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x00, 0x56, 0xA6)
    p_just = doc.add_paragraph(justificativa_gestor)
    p_just.paragraph_format.line_spacing = 1.15
    p_just.paragraph_format.space_after = Pt(12)

    # 3. Tabela Orçamental
    h3 = doc.add_paragraph()
    r3 = h3.add_run("3. DISTRIBUIÇÃO ORÇAMENTÁRIA POR CONTA CONTÁBIL")
    r3.bold = True
    r3.font.size = Pt(12)
    r3.font.color.rgb = RGBColor(0x00, 0x56, 0xA6)

    df_tb = resumo_contas.reset_index()
    table_fin = doc.add_table(rows=len(df_tb) + 2, cols=len(df_tb.columns))
    table_fin.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Cabeçalho Tabela
    hdr_cells = table_fin.rows[0].cells
    for c_idx, col_name in enumerate(df_tb.columns):
        hdr_cells[c_idx].paragraphs[0].add_run(str(col_name)).bold = True
        hdr_cells[c_idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[c_idx], '0056A6')
        set_cell_margins(hdr_cells[c_idx], top=80, bottom=80)

    # Linhas de Dados
    for r_idx, row_data in df_tb.iterrows():
        row_cells = table_fin.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            cell_p = row_cells[c_idx].paragraphs[0]
            if isinstance(val, (int, float)) and c_idx > 0:
                cell_p.add_run(f"R$ {val:,.2f}")
                cell_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                cell_p.add_run(str(val))
            set_cell_margins(row_cells[c_idx], top=60, bottom=60)
            if r_idx % 2 == 1:
                set_cell_background(row_cells[c_idx], 'F9FAFB')

    # Totais
    tot_cells = table_fin.rows[-1].cells
    tot_cells[0].paragraphs[0].add_run("TOTAL GERAL").bold = True
    set_cell_background(tot_cells[0], 'E5E7EB')

    if 'CAPEX' in df_tb.columns:
        c_idx = df_tb.columns.get_loc('CAPEX')
        p = tot_cells[c_idx].paragraphs[0]
        p.add_run(f"R$ {total_capex:,.2f}").bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_cell_background(tot_cells[c_idx], 'E5E7EB')

    if 'OPEX' in df_tb.columns:
        c_idx = df_tb.columns.get_loc('OPEX')
        p = tot_cells[c_idx].paragraphs[0]
        p.add_run(f"R$ {total_opex:,.2f}").bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_cell_background(tot_cells[c_idx], 'E5E7EB')

    if 'Total Geral' in df_tb.columns:
        c_idx = df_tb.columns.get_loc('Total Geral')
        p = tot_cells[c_idx].paragraphs[0]
        p.add_run(f"R$ {total_geral:,.2f}").bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_cell_background(tot_cells[c_idx], 'E5E7EB')

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Observações
    if observacoes_usuario.strip():
        h4 = doc.add_paragraph()
        r4 = h4.add_run("4. OBSERVAÇÕES COMPLEMENTARES")
        r4.bold = True
        r4.font.size = Pt(12)
        r4.font.color.rgb = RGBColor(0x00, 0x56, 0xA6)
        p_obs = doc.add_paragraph(observacoes_usuario)
        p_obs.paragraph_format.line_spacing = 1.15

    # Retorna o arquivo em memória para download
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# -----------------------------------------------------------------------------
# 4. INTERFACE DO UTILIZADOR (TABS DO STREAMLIT)
# -----------------------------------------------------------------------------
tab1, tab2 = st.tabs(["📄 1. Importação & Dados da RFA", "📊 2. Resumo & Exportação Word"])

with tab1:
    col_left, col_right = st.columns([1, 1])
    
    with col_left:
        st.subheader("1.1 Carregar Caderno / Planilha")
        uploaded_file = st.file_uploader("Selecione a planilha (.xlsx)", type=["xlsx"])
        
        st.subheader("1.2 Identificação e Solicitante")
        gerente_resp = st.text_input("Gerente Responsável", value="Adriana Paula Cassol")
        unidade_regional = st.text_input("Regional / Unidade", value="SENAI/SC - Regional Oeste - SENAI/SC - Chapecó")
        descricao_sucinta = st.text_input("Descrição Sucinta (Título)", value="Aquisição de Móveis e Equipamentos para Biblioteca e 3º Pavimento do SENAI Chapecó")
        periodo_desembolso = st.text_input("Período de Desembolso Previsto", value="jan/27 a mar/27")

    with col_right:
        st.subheader("1.3 Justificativa & Observações")
        justificativa_gestor = st.text_area(
            "Justificativa (Viabilidade / Payback / Manutenção / Risco)",
            height=160,
            value="Com a conclusão e entrega da obra do IST Alimentos e Bebidas em Chapecó prevista para o segundo semestre de 2026, faz-se necessária a aquisição de acervo bibliográfico, mobiliário e equipamentos para a implantação da Biblioteca no 2º pavimento da unidade e equipamentos e mobiliário para 10 novos laboratórios no 3º pavimento. Os espaços atenderão Educação do SENAI em todas as suas modalidades. O investimento proposto é fundamental para garantir a conformidade regulatória dos cursos e a prontidão operacional do novo ambiente de ensino já no 1º semestre de 2027."
        )
        observacoes_usuario = st.text_area(
            "Observações Complementares",
            height=100,
            value="- Reduzido os itens de mobiliários e TI já existentes que serão transferidos.\n- Adicionamos 10% nos valores, exceto equipamentos de informática.\n- O valor das contas pode sofrer alterações devido à economicidade de aquisição."
        )

# -----------------------------------------------------------------------------
# 5. PROCESSAMENTO E EXIBIÇÃO NA TAB 2
# -----------------------------------------------------------------------------
if uploaded_file is not None:
    df_proc = processar_planilha(uploaded_file)
    
    # Agrupamento
    col_conta = 'ITEM' if 'ITEM' in df_proc.columns else df_proc.columns[0]
    
    if 'TIPO_DESPESA' in df_proc.columns:
        resumo_contas = df_proc.groupby([col_conta, 'TIPO_DESPESA'])['VALOR_TOTAL'].sum().unstack(fill_value=0)
    else:
        resumo_contas = df_proc.groupby(col_conta)['VALOR_TOTAL'].sum().to_frame()
        resumo_contas.columns = ['CAPEX']

    if 'CAPEX' not in resumo_contas.columns: resumo_contas['CAPEX'] = 0.0
    if 'OPEX' not in resumo_contas.columns: resumo_contas['OPEX'] = 0.0
    resumo_contas['Total Geral'] = resumo_contas['CAPEX'] + resumo_contas['OPEX']

    total_capex = resumo_contas['CAPEX'].sum()
    total_opex = resumo_contas['OPEX'].sum()
    total_geral = resumo_contas['Total Geral'].sum()

    # Identificação dos Chamados
    c_getic = df_proc['GETIC_CHAMADO'].dropna().unique() if 'GETIC_CHAMADO' in df_proc.columns else ['TI0364194']
    c_genge = df_proc['GENGE_CHAMADO'].dropna().unique() if 'GENGE_CHAMADO' in df_proc.columns else ['N/A']
    chamados_getic_str = ', '.join([str(c) for c in c_getic if str(c) not in ['nan', 'None']])
    chamados_genge_str = ', '.join([str(c) for c in c_genge if str(c) not in ['nan', 'None']])

    with tab2:
        st.subheader("📈 Resumo Financeiro da RFA")
        kpi1, kpi2, kpi3 = st.columns(3)
        kpi1.metric("CAPEX (Investimento)", f"R$ {total_capex:,.2f}")
        kpi2.metric("OPEX (Custeio)", f"R$ {total_opex:,.2f}")
        kpi3.metric("VALOR TOTAL RFA", f"R$ {total_geral:,.2f}")
        
        st.divider()
        st.subheader("📋 Quadro por Conta Contábil")
        st.dataframe(resumo_contas.style.format("R$ {:,.2f}"), use_container_width=True)
        
        st.divider()
        # Geração do arquivo em memória
        docx_buffer = gerar_docx_rfa_bytes(
            gerente_resp, unidade_regional, descricao_sucinta, periodo_desembolso,
            justificativa_gestor, observacoes_usuario, resumo_contas,
            total_capex, total_opex, total_geral, chamados_getic_str, chamados_genge_str
        )
        
        st.download_button(
            label="📥 Baixar Minuta da RFA Formatada em Word (.docx)",
            data=docx_buffer,
            file_name="RFA_SE_Suite_Chapeco.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
else:
    with tab2:
        st.info("👈 Por favor, carregue a planilha (.xlsx) na Aba 1 para visualizar o resumo financeiro e gerar o documento.")
